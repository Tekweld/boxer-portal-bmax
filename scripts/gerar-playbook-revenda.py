# Gera Documentação/Playbook_Revenda_Portal_BMAX.docx DO ZERO (python-docx) — 2 páginas A4.
# Regra do projeto: DOCX é sempre regenerado, nunca editado no XML. Atualize os textos aqui e rode:
#   python -I scripts/gerar-playbook-revenda.py
# Estilo: padrão Boxer (Arial, azul #1D327B, ciano #25BBEE, vermelho #E30613).
# Revisão 14/10/2026: Nova Negociação + trava de CNPJ, prazos (48h úteis / 60 dias corridos), % de comissão
# SÓ dos perfis em que a revenda participa (nada de comissão de representante), perguntas frequentes.
# Os percentuais vêm da tabela comissao_tabela (Supabase boxer-sistemas) — se mudarem, atualize COMISSOES abaixo.
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor

SAIDA = Path(__file__).resolve().parent.parent.parent / "Documentação" / "Playbook_Revenda_Portal_BMAX.docx"

AZUL, CIANO, VERMELHO = "1D327B", "25BBEE", "E30613"
TEXTO, SEC, FUNDO = "1A202C", "4A5568", "F0F4F8"
TAM = 9.5  # corpo

# % de comissão da REVENDA por classe de preço (1..6) — somente perfis em que ela participa
COMISSOES = [
    ("Lead enviado pela Boxer · Laser · cliente a até 50 km da sua revenda", "PCI 6", [1.5] * 6),
    ("Lead enviado pela Boxer · Máquina industrial acima de R$ 25 mil · até 50 km", "PCI 9", [2] * 6),
    ("Lead enviado pela Boxer · Máquina industrial até R$ 25 mil · até 50 km · você escolhe “Boxer vende”", "PCI 12", [2] * 6),
    ("Cliente trazido por você · Robô · Boxer vende", "PCI 13", [1.5, 2, 2.5, 3, 3.5, 4]),
    ("Cliente trazido por você · Laser · Boxer vende", "PCI 14", [1.7, 2, 2.5, 3, 3.5, 4]),
    ("Cliente trazido por você · Máquina industrial · Boxer vende", "PCI 15", [3, 3.5, 4, 4.5, 5, 5]),
]

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(1.9)
sec.top_margin, sec.bottom_margin = Cm(1.3), Cm(1.4)
sec.footer_distance = Cm(0.6)

normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
normal.font.size = Pt(TAM)


def run(p, texto, tam=TAM, cor=TEXTO, negrito=False, italico=False):
    r = p.add_run(texto)
    r.font.name = "Arial"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    r.font.size = Pt(tam)
    r.font.color.rgb = RGBColor.from_string(cor)
    r.bold = negrito or None
    r.italic = italico or None
    return r


def espaco(p, antes=0, depois=0, linha=1.08):
    p.paragraph_format.space_before = Pt(antes)
    p.paragraph_format.space_after = Pt(depois)
    p.paragraph_format.line_spacing = linha


def rico(p, partes, tam=TAM, cor=TEXTO):
    for texto, neg in partes:
        run(p, texto, tam=tam, negrito=neg, cor=cor)


# ── Numeração ───────────────────────────────────────────────────────────────────────
numbering = doc.part.numbering_part.element
numbering.append(parse_xml(
    f'<w:abstractNum {nsdecls("w")} w:abstractNumId="90"><w:multiLevelType w:val="hybridMultilevel"/>'
    '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/><w:lvlJc w:val="left"/>'
    '<w:pPr><w:ind w:left="400" w:hanging="300"/></w:pPr>'
    f'<w:rPr><w:rFonts w:ascii="Arial" w:hAnsi="Arial" w:cs="Arial"/><w:b/><w:color w:val="{AZUL}"/></w:rPr></w:lvl></w:abstractNum>'))
numbering.append(parse_xml(
    f'<w:abstractNum {nsdecls("w")} w:abstractNumId="91"><w:multiLevelType w:val="hybridMultilevel"/>'
    '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="•"/><w:lvlJc w:val="left"/>'
    '<w:pPr><w:ind w:left="400" w:hanging="230"/></w:pPr>'
    f'<w:rPr><w:rFonts w:ascii="Arial" w:hAnsi="Arial" w:cs="Arial"/><w:color w:val="{AZUL}"/></w:rPr></w:lvl></w:abstractNum>'))
numbering.append(parse_xml(f'<w:num {nsdecls("w")} w:numId="90"><w:abstractNumId w:val="90"/></w:num>'))
numbering.append(parse_xml(f'<w:num {nsdecls("w")} w:numId="91"><w:abstractNumId w:val="91"/></w:num>'))
nums = [e for e in numbering if e.tag == qn("w:num")]
abstratos = [e for e in numbering if e.tag == qn("w:abstractNum")]
for e in abstratos + nums:
    numbering.remove(e)
for e in abstratos + nums:
    numbering.append(e)


def item(partes, num_id, depois=2.5):
    p = doc.add_paragraph()
    p._p.get_or_add_pPr().append(parse_xml(
        f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/><w:numId w:val="{num_id}"/></w:numPr>'))
    espaco(p, 0, depois)
    rico(p, partes)
    return p


def secao(titulo, antes=7):
    p = doc.add_paragraph()
    espaco(p, antes, 3)
    p.paragraph_format.keep_with_next = True
    p._p.get_or_add_pPr().append(parse_xml(
        f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="8" w:space="1" w:color="{CIANO}"/></w:pBdr>'))
    run(p, titulo, tam=12, cor=AZUL, negrito=True)


def faixa(p):
    pPr = p._p.get_or_add_pPr()
    pPr.append(parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="24" w:space="3" w:color="{AZUL}"/>'
        f'<w:left w:val="single" w:sz="24" w:space="4" w:color="{AZUL}"/>'
        f'<w:bottom w:val="single" w:sz="24" w:space="3" w:color="{AZUL}"/>'
        f'<w:right w:val="single" w:sz="24" w:space="4" w:color="{AZUL}"/></w:pBdr>'))
    pPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{AZUL}"/>'))


def celula(cel, texto, tam=8.5, cor=TEXTO, negrito=False, fundo=None, centro=False, largura_cm=None):
    if largura_cm:
        cel.width = Cm(largura_cm)
    tcPr = cel._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(
        f'<w:tcMar {nsdecls("w")}><w:top w:w="40" w:type="dxa"/><w:left w:w="80" w:type="dxa"/>'
        '<w:bottom w:w="40" w:type="dxa"/><w:right w:w="80" w:type="dxa"/></w:tcMar>'))
    if fundo:
        tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{fundo}"/>'))
    p = cel.paragraphs[0]
    espaco(p, 0, 0, 1.0)
    if centro:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, texto, tam=tam, cor=cor, negrito=negrito)


def tabela(larguras_cm, cabecalho, linhas, centro_a_partir=1):
    t = doc.add_table(rows=1 + len(linhas), cols=len(larguras_cm))
    t.autofit = False
    t.alignment = 0
    t._tbl.tblPr.append(parse_xml(
        f'<w:tblBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:color="D0D8E8"/>'
        '<w:left w:val="single" w:sz="4" w:color="D0D8E8"/><w:bottom w:val="single" w:sz="4" w:color="D0D8E8"/>'
        '<w:right w:val="single" w:sz="4" w:color="D0D8E8"/><w:insideH w:val="single" w:sz="4" w:color="D0D8E8"/>'
        '<w:insideV w:val="single" w:sz="4" w:color="D0D8E8"/></w:tblBorders>'))
    for j, h in enumerate(cabecalho):
        celula(t.rows[0].cells[j], h, tam=8, cor="FFFFFF", negrito=True, fundo=AZUL, centro=j >= centro_a_partir, largura_cm=larguras_cm[j])
    for i, linha in enumerate(linhas, start=1):
        for j, v in enumerate(linha):
            celula(t.rows[i].cells[j], v, negrito=(j == 0 and len(cabecalho) <= 3), fundo=(FUNDO if i % 2 == 0 else None),
                   centro=j >= centro_a_partir, largura_cm=larguras_cm[j])
    return t


def pct(v):
    return (f"{v:g}".replace(".", ",")) + "%"


# ═══ PÁGINA 1 ═══════════════════════════════════════════════════════════════════════
p = doc.add_paragraph(); espaco(p, 4, 0); faixa(p)
run(p, "  PLAYBOOK DA REVENDA", tam=19, cor="FFFFFF", negrito=True)
p = doc.add_paragraph(); espaco(p, 0, 6); faixa(p)
run(p, "  Como usar o Portal BMAX no dia a dia", tam=11, cor="BFD7F5")

p = doc.add_paragraph(); espaco(p, 4, 0)
run(p, "Acesse em ")
run(p, "bmax.boxersoldas.com.br", tam=11, cor=AZUL, negrito=True)
run(p, "  •  login pessoal de cada colaborador", cor=SEC)

secao("Sua rotina em 5 passos", antes=5)
item([("Abra o painel. ", True), ("Confira os leads novos e os que estão em andamento na sua revenda.", False)], 90)
item([("Defina como atender cada lead enviado pela Boxer. ", True),
      ("No card, escolha em “Como deseja atender este lead?”: ", False), ("Eu assumo a venda", True), (" ou ", False),
      ("Boxer vende, eu ganho comissão", True), (". Você tem ", False), ("48 horas úteis", True),
      (" (sábado e domingo não contam) e recebe lembrete por e-mail. Sem resposta, o lead segue como “Boxer vende” e o time Boxer assume — você mantém a comissão.", False)], 90)
item([("Se você assumiu, atenda e registre o resultado. ", True),
      ("Os dados de contato do cliente passam a aparecer no card. Ao concluir, marque ", False), ("Vendido", True),
      (" (com o valor real da venda) ou ", False), ("Perdido", True), (" em até ", False), ("60 dias corridos", True),
      (" desde que assumiu. O resultado é definitivo: marque só o desfecho real.", False)], 90)
item([("Se a Boxer vai vender, acompanhe. ", True),
      ("Você vê o andamento no card e não precisa fazer mais nada: a comissão é apurada quando a venda se efetiva.", False)], 90)
item([("Trouxe um cliente novo? ", True), ("Use ", False), ("+ Nova Negociação", True), (" (explicada abaixo).", False)], 90)

secao("Nova Negociação: registrando um cliente que você trouxe")
item([("O que preencher: ", True),
      ("CNPJ, nome do cliente, cidade, estado, produto, caminho de venda e ", False), ("orçamento", True),
      (" são obrigatórios. CEP, máquina de interesse e observações são opcionais. Em ", False), ("Orçamento", True),
      (", anexe de ", False), ("1 a 3 arquivos", True),
      (" (PDF, imagem, Word ou Excel, até 10 MB cada): o time Boxer recebe o link no card do RD, sem precisar de login e sem prazo para expirar.", False)], 91)
item([("Produto: ", True), ("Máquinas, Laser ou Robô. Ele define o PCI do lead (15, 14 ou 13, tabela da página 2) e não muda com o caminho escolhido. ", False),
      ("Caminho de venda: ", True), ("Eu assumo a venda", True),
      (" → a venda é sua e o cliente fica registrado em nome da sua revenda; passam a contar os ", False),
      ("60 dias corridos", True), (" para registrar Vendido ou Perdido. ", False), ("Boxer vende, eu ganho comissão", True),
      (" → a Boxer fatura e você recebe o cashback. Nesse caso, diga ", False), ("quem atende o cliente", True),
      (": se for ", False), ("você", True), (", o lead não aparece para vendedor da Boxer; se for ", False), ("um vendedor Boxer", True),
      (", ele chega ao vendedor responsável da sua região.", False)], 91)
item([("Por que meu cadastro foi recusado? ", True),
      ("Antes de salvar, o Portal confere se esse CNPJ já tem uma negociação ativa na Boxer (de qualquer revenda). Se tiver, aparece um aviso vermelho — "
       "“Cadastro NÃO foi salvo” — com o nome do lead que já existe, e nada é criado. Cada cliente só pode ter uma negociação ativa por vez, "
       "o que protege quem registrou primeiro. Se achar que é engano, fale com o seu representante.", False)], 91)
item([("Cadastre só negociações reais. ", True),
      ("Ao assumir, o CNPJ fica reservado para a sua revenda por 60 dias: nesse período nenhuma outra revenda (nem a Boxer) consegue cadastrar o mesmo cliente. "
       "Se não houver Vendido ou Perdido até o 60º dia, o lead é arquivado como “Prazo Vencido” e o CNPJ volta a ficar livre para qualquer revenda. "
       "Cadastrar só para “garantir” o cliente, sem negociar de verdade, trava o CNPJ sem gerar venda.", False)], 91)

secao("Prazos em um olhar")
tabela([5.9, 3.2, 8.1], ["Prazo", "Quanto", "O que acontece"], [
    ["Definir como atender o lead enviado pela Boxer", "48 horas úteis", "Sábado e domingo não contam. Sem resposta, o lead segue como “Boxer vende”."],
    ["Registrar Vendido ou Perdido após assumir", "60 dias corridos", "Sem resultado, o lead é arquivado como “Prazo Vencido” e o CNPJ é liberado."],
    ["Validade do cashback", "6 meses", "O extrato e o e-mail avisam o que está perto de vencer."],
    ["Cheque de cashback aprovado", "30 dias", "Informe o código no campo Observação do pedido."],
], centro_a_partir=99)

# ═══ PÁGINA 2 ═══════════════════════════════════════════════════════════════════════
doc.add_page_break()
secao("Quanto você ganha (cashback)", antes=0)
p = doc.add_paragraph(); espaco(p, 0, 3)
rico(p, [("Seu cashback é um percentual sobre o ", False), ("valor real da venda", True),
         (", creditado quando a venda se efetiva (Vendido). Abaixo, só os perfis em que a sua revenda participa. "
          "A sua ", False), ("classe de preço", True), (" está no cadastro BMAX — confirme com o seu representante.", False)])
tabela([6.9, 1.4] + [1.45] * 6, ["Situação do lead", "PCI"] + [f"Classe {i}" for i in range(1, 7)],
       [[d, pci] + [pct(v) for v in vals] for d, pci, vals in COMISSOES], centro_a_partir=1)
p = doc.add_paragraph(); espaco(p, 3, 0)
run(p, "Em lead enviado pela Boxer, se você escolher “Eu assumo” (PCI 12a), a venda é sua e não há comissão da Boxer sobre ela.", tam=8.5, cor=SEC, italico=True)

secao("Como acompanhar e usar o cashback")
item([("Extrato: ", True), ("clique em ", False), ("ver extrato", True),
      (" no card “Cashback Total”: saldo disponível e cada movimentação, com o nome do cliente.", False)], 91)
item([("Como usar: ", True), ("no extrato, ", False), ("Solicitar Saque", True), (" → valor e tipo de uso (", False), ("Desconto", True),
      (" ou ", False), ("Bonificação", True), (", com os limites exibidos na tela) → aprovação do seu representante → você recebe o cheque com um código (também em PDF).", False)], 91)
item([("Vale para ", True), ("compra de máquinas do portfólio Boxer, não para peças e consumíveis.", False)], 91)

secao("E se...?")
item([("Esqueci minha senha: ", True),
      ("avise o seu representante Boxer. Ele pede a redefinição ao administrador e a nova senha chega por canal seguro (ex.: WhatsApp). "
       "Nunca compartilhe a senha — cada colaborador tem o próprio login; para acesso de um novo colaborador, peça ao seu representante.", False)], 91)
item([("Meu lead não aparece no painel: ", True),
      ("só equipamentos BMAX entram no programa (máquinas, robôs e laser). Consumíveis, acessórios avulsos e outros produtos não aparecem e não geram cashback. "
       "Escreva o equipamento pelo nome da tabela de preços BMAX (ex.: Hardmig 450, LQ 1250); nome incompleto ou de outra linha pode deixar o card de fora.", False)], 91)
item([("Escolhi o caminho errado: ", True),
      ("o Portal não desfaz a escolha. Fale com o seu representante o quanto antes.", False)], 91)
item([("Meu prazo de 60 dias venceu: ", True),
      ("o lead foi arquivado como “Prazo Vencido” e o CNPJ está livre. Se o cliente retomou a negociação, registre uma nova em ", False),
      ("+ Nova Negociação", True), (".", False)], 91)
item([("O cashback não bateu: ", True),
      ("o valor vem do valor real informado em Vendido e da tabela acima. Confira cidade, equipamento e valor no card e fale com o seu representante.", False)], 91)

# ── Confidencialidade ───────────────────────────────────────────────────────────────
p = doc.add_paragraph(); espaco(p, 4, 0)
caixa = doc.add_table(rows=1, cols=1)
caixa.autofit = False
caixa._tbl.tblPr.append(parse_xml(f'<w:tblW {nsdecls("w")} w:type="dxa" w:w="9640"/>'))
cel = caixa.rows[0].cells[0]
cel.width = Cm(17.0)
tcPr = cel._tc.get_or_add_tcPr()
tcPr.append(parse_xml(
    f'<w:tcBorders {nsdecls("w")}><w:top w:val="nil"/><w:left w:val="single" w:sz="24" w:color="{VERMELHO}"/>'
    '<w:bottom w:val="nil"/><w:right w:val="nil"/></w:tcBorders>'))
tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{FUNDO}"/>'))
tcPr.append(parse_xml(
    f'<w:tcMar {nsdecls("w")}><w:top w:w="100" w:type="dxa"/><w:left w:w="180" w:type="dxa"/>'
    '<w:bottom w:w="100" w:type="dxa"/><w:right w:w="180" w:type="dxa"/></w:tcMar>'))
p = cel.paragraphs[0]; espaco(p, 0, 2)
run(p, "Confidencialidade", cor=VERMELHO, negrito=True)
for t in ("• Use apenas o seu login e não compartilhe senha. Novo colaborador? Peça o acesso ao seu representante Boxer.",
          "• Dados de clientes, condições e valores do programa são restritos à sua revenda. Não divulgue este material nem prints do portal a terceiros."):
    p = cel.add_paragraph(); espaco(p, 0, 1.5)
    run(p, t, tam=9, cor=SEC)

p = doc.add_paragraph(); espaco(p, 5, 0)
run(p, "Dúvidas ou algo estranho no painel? Fale com o seu representante Boxer.", tam=9.5, cor=SEC, italico=True)

# ── Rodapé com número de página ─────────────────────────────────────────────────────
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
run(fp, "Uso exclusivo das revendas parceiras BMAX • não compartilhar  |  Página ", tam=8, cor=SEC)
r = run(fp, "", tam=8, cor=SEC)
for tipo, texto in (("begin", None), (None, "PAGE"), ("end", None)):
    if tipo:
        r._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="{tipo}"/>'))
    else:
        r._r.append(parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve">{texto}</w:instrText>'))

SAIDA.parent.mkdir(exist_ok=True)
doc.save(SAIDA)
print("gerado:", SAIDA)
